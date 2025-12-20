import sys
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QTextEdit, QFileDialog,
    QToolBar, QFontComboBox, QComboBox, QMessageBox
)
from PyQt6.QtGui import QAction, QFont, QTextCursor
from PyQt6.QtCore import Qt
from docx import Document
from docx.shared import Pt


class WordCloneApp(QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Word Clone - PyQt6")
        self.setGeometry(100, 100, 1000, 700)

        self.file_path = None

        self.editor = QTextEdit()
        self.setCentralWidget(self.editor)

        self.create_menu()
        self.create_toolbar()
        self.create_shortcuts()

    # =========================
    # MENU
    # =========================
    def create_menu(self):
        menu_bar = self.menuBar()

        file_menu = menu_bar.addMenu("Fichier")
        edit_menu = menu_bar.addMenu("Édition")
        format_menu = menu_bar.addMenu("Format")

        file_menu.addAction("Nouveau", self.new_file)
        file_menu.addAction("Ouvrir", self.open_file)
        file_menu.addAction("Sauvegarder", self.save_file)
        file_menu.addSeparator()
        file_menu.addAction("Quitter", self.close)

        edit_menu.addAction("Gras", self.toggle_bold)
        edit_menu.addAction("Italique", self.toggle_italic)
        edit_menu.addAction("Souligné", self.toggle_underline)

        format_menu.addAction("Titre H1", lambda: self.insert_heading(1))
        format_menu.addAction("Titre H2", lambda: self.insert_heading(2))
        format_menu.addAction("Liste à puces", self.insert_bullet)
        format_menu.addAction("Liste numérotée", self.insert_numbered)

    # =========================
    # BARRE D’OUTILS
    # =========================
    def create_toolbar(self):
        toolbar = QToolBar("Outils")
        self.addToolBar(toolbar)

        # Actions
        new_action = QAction("Nouveau", self)
        open_action = QAction("Ouvrir", self)
        save_action = QAction("Sauvegarder", self)

        bold_action = QAction("B", self)
        italic_action = QAction("I", self)
        underline_action = QAction("U", self)

        bold_action.triggered.connect(self.toggle_bold)
        italic_action.triggered.connect(self.toggle_italic)
        underline_action.triggered.connect(self.toggle_underline)

        new_action.triggered.connect(self.new_file)
        open_action.triggered.connect(self.open_file)
        save_action.triggered.connect(self.save_file)

        toolbar.addAction(new_action)
        toolbar.addAction(open_action)
        toolbar.addAction(save_action)
        toolbar.addSeparator()
        toolbar.addAction(bold_action)
        toolbar.addAction(italic_action)
        toolbar.addAction(underline_action)

        # Police
        self.font_box = QFontComboBox()
        self.font_box.currentFontChanged.connect(self.change_font)
        toolbar.addWidget(self.font_box)

        # Taille police
        self.font_size_box = QComboBox()
        self.font_size_box.addItems(
            [str(size) for size in [10, 12, 14, 16, 18, 20, 24, 28]]
        )
        self.font_size_box.setCurrentText("12")
        self.font_size_box.currentTextChanged.connect(self.change_font_size)
        toolbar.addWidget(self.font_size_box)

    # =========================
    # RACCOURCIS
    # =========================
    def create_shortcuts(self):
        QAction("Gras", self, shortcut="Ctrl+B", triggered=self.toggle_bold)
        QAction("Italique", self, shortcut="Ctrl+I", triggered=self.toggle_italic)
        QAction("Souligné", self, shortcut="Ctrl+U", triggered=self.toggle_underline)

    # =========================
    # FICHIERS
    # =========================
    def new_file(self):
        self.editor.clear()
        self.file_path = None

    def open_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Ouvrir un fichier", "", "Word (*.docx)"
        )
        if not path:
            return

        doc = Document(path)
        self.editor.clear()

        for para in doc.paragraphs:
            self.editor.append(para.text)

        self.file_path = path

    def save_file(self):
        if not self.file_path:
            self.file_path, _ = QFileDialog.getSaveFileName(
                self, "Sauvegarder", "", "Word (*.docx)"
            )

        if not self.file_path:
            return

        doc = Document()
        for line in self.editor.toPlainText().split("\n"):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(int(self.font_size_box.currentText()))

        doc.save(self.file_path)
        QMessageBox.information(self, "Succès", "Document sauvegardé")

    # =========================
    # FORMATAGE TEXTE
    # =========================
    def toggle_bold(self):
        fmt = self.editor.currentCharFormat()
        fmt.setFontWeight(
            QFont.Weight.Bold
            if fmt.fontWeight() != QFont.Weight.Bold
            else QFont.Weight.Normal
        )
        self.editor.setCurrentCharFormat(fmt)

    def toggle_italic(self):
        fmt = self.editor.currentCharFormat()
        fmt.setFontItalic(not fmt.fontItalic())
        self.editor.setCurrentCharFormat(fmt)

    def toggle_underline(self):
        fmt = self.editor.currentCharFormat()
        fmt.setFontUnderline(not fmt.fontUnderline())
        self.editor.setCurrentCharFormat(fmt)

    def change_font(self, font):
        fmt = self.editor.currentCharFormat()
        fmt.setFontFamily(font.family())
        self.editor.setCurrentCharFormat(fmt)

    def change_font_size(self, size):
        fmt = self.editor.currentCharFormat()
        fmt.setFontPointSize(float(size))
        self.editor.setCurrentCharFormat(fmt)

    # =========================
    # TITRES & LISTES
    # =========================
    def insert_heading(self, level):
        cursor = self.editor.textCursor()
        font = QFont("Arial", 24 if level == 1 else 18)
        font.setBold(True)
        cursor.insertText("\nTitre\n", font)

    def insert_bullet(self):
        self.editor.insertPlainText("• Élément de liste\n")

    def insert_numbered(self):
        self.editor.insertPlainText("1. Élément numéroté\n")


# =========================
# MAIN
# =========================
if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = WordCloneApp()
    window.show()
    sys.exit(app.exec())
