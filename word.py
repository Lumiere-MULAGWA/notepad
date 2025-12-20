import tkinter as tk
from tkinter import filedialog, messagebox, font
from tkinter.ttk import Combobox
from docx import Document
from docx.shared import Pt


class WordCloneApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Word Clone - Python")
        self.root.geometry("900x600")

        self.file_path = None

        self.create_menu()
        self.create_toolbar()
        self.create_text_area()
        self.create_shortcuts()

    # =========================
    # MENU
    # =========================
    def create_menu(self):
        menu_bar = tk.Menu(self.root)

        # Fichier
        file_menu = tk.Menu(menu_bar, tearoff=0)
        file_menu.add_command(label="Nouveau", command=self.new_file)
        file_menu.add_command(label="Ouvrir", command=self.open_file)
        file_menu.add_command(label="Sauvegarder", command=self.save_file)
        file_menu.add_separator()
        file_menu.add_command(label="Quitter", command=self.root.quit)

        # Édition
        edit_menu = tk.Menu(menu_bar, tearoff=0)
        edit_menu.add_command(label="Gras", command=self.toggle_bold)
        edit_menu.add_command(label="Italique", command=self.toggle_italic)
        edit_menu.add_command(label="Souligné", command=self.toggle_underline)

        # Format
        format_menu = tk.Menu(menu_bar, tearoff=0)
        format_menu.add_command(label="Titre H1", command=lambda: self.insert_heading(1))
        format_menu.add_command(label="Titre H2", command=lambda: self.insert_heading(2))
        format_menu.add_command(label="Liste à puces", command=self.insert_bullet)
        format_menu.add_command(label="Liste numérotée", command=self.insert_numbered)

        menu_bar.add_cascade(label="Fichier", menu=file_menu)
        menu_bar.add_cascade(label="Édition", menu=edit_menu)
        menu_bar.add_cascade(label="Format", menu=format_menu)

        self.root.config(menu=menu_bar)

    # =========================
    # BARRE D’OUTILS
    # =========================
    def create_toolbar(self):
        toolbar = tk.Frame(self.root, bd=1, relief=tk.RAISED)

        tk.Button(toolbar, text="Nouveau", command=self.new_file).pack(side=tk.LEFT)
        tk.Button(toolbar, text="Ouvrir", command=self.open_file).pack(side=tk.LEFT)
        tk.Button(toolbar, text="Sauvegarder", command=self.save_file).pack(side=tk.LEFT)

        tk.Button(toolbar, text="B", command=self.toggle_bold).pack(side=tk.LEFT)
        tk.Button(toolbar, text="I", command=self.toggle_italic).pack(side=tk.LEFT)
        tk.Button(toolbar, text="U", command=self.toggle_underline).pack(side=tk.LEFT)

        # Police
        self.font_family = Combobox(
            toolbar, values=font.families(), width=20
        )
        self.font_family.set("Arial")
        self.font_family.bind("<<ComboboxSelected>>", self.change_font)
        self.font_family.pack(side=tk.LEFT)

        # Taille
        self.font_size = Combobox(
            toolbar, values=[10, 12, 14, 16, 18, 20, 24, 28], width=5
        )
        self.font_size.set(12)
        self.font_size.bind("<<ComboboxSelected>>", self.change_font)
        self.font_size.pack(side=tk.LEFT)

        toolbar.pack(fill=tk.X)

    # =========================
    # ZONE DE TEXTE
    # =========================
    def create_text_area(self):
        self.text_area = tk.Text(self.root, wrap="word", undo=True)
        self.text_area.pack(expand=1, fill="both")

        self.default_font = font.Font(
            family="Arial", size=12
        )
        self.text_area.configure(font=self.default_font)

        # Tags
        self.text_area.tag_configure("bold", font=font.Font(weight="bold"))
        self.text_area.tag_configure("italic", font=font.Font(slant="italic"))
        self.text_area.tag_configure("underline", font=font.Font(underline=1))

    # =========================
    # RACCOURCIS CLAVIER
    # =========================
    def create_shortcuts(self):
        self.root.bind("<Control-b>", lambda e: self.toggle_bold())
        self.root.bind("<Control-i>", lambda e: self.toggle_italic())
        self.root.bind("<Control-u>", lambda e: self.toggle_underline())

    # =========================
    # FONCTIONS FICHIER
    # =========================
    def new_file(self):
        self.text_area.delete(1.0, tk.END)
        self.file_path = None

    def open_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Word Documents", "*.docx")]
        )
        if not path:
            return

        doc = Document(path)
        self.text_area.delete(1.0, tk.END)

        for para in doc.paragraphs:
            self.text_area.insert(tk.END, para.text + "\n")

        self.file_path = path

    def save_file(self):
        if not self.file_path:
            self.file_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word Documents", "*.docx")]
            )

        if not self.file_path:
            return

        doc = Document()
        content = self.text_area.get(1.0, tk.END).split("\n")

        for line in content:
            p = doc.add_paragraph(line)
            for run in p.runs:
                run.font.size = Pt(int(self.font_size.get()))

        doc.save(self.file_path)
        messagebox.showinfo("Succès", "Document sauvegardé avec succès")

    # =========================
    # FORMATAGE TEXTE
    # =========================
    def toggle_tag(self, tag):
        try:
            start, end = self.text_area.tag_ranges(tk.SEL)
            if tag in self.text_area.tag_names("sel.first"):
                self.text_area.tag_remove(tag, start, end)
            else:
                self.text_area.tag_add(tag, start, end)
        except ValueError:
            pass

    def toggle_bold(self):
        self.toggle_tag("bold")

    def toggle_italic(self):
        self.toggle_tag("italic")

    def toggle_underline(self):
        self.toggle_tag("underline")

    def change_font(self, event=None):
        self.default_font.config(
            family=self.font_family.get(),
            size=int(self.font_size.get())
        )

    # =========================
    # TITRES & LISTES
    # =========================
    def insert_heading(self, level):
        self.text_area.insert(
            tk.INSERT, f"\n[Titre H{level}]\n"
        )

    def insert_bullet(self):
        self.text_area.insert(tk.INSERT, "• Élément de liste\n")

    def insert_numbered(self):
        self.text_area.insert(tk.INSERT, "1. Élément numéroté\n")


# =========================
# LANCEMENT APPLICATION
# =========================
if __name__ == "__main__":
    root = tk.Tk()
    app = WordCloneApp(root)
    root.mainloop()
